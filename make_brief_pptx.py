"""Short Agentic AI briefing deck (~12 slides)."""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

PAPER = RGBColor(0xF1, 0xF1, 0xF1)
CREAM = RGBColor(0xDA, 0xCE, 0xBE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1A, 0x16, 0x12)
MUTED = RGBColor(0x5E, 0x57, 0x4E)
ORANGE = RGBColor(0xDB, 0x8F, 0x2A)
ORANGE_DARK = RGBColor(0xB8, 0x74, 0x1C)
LINE = RGBColor(0xD4, 0xC8, 0xB8)

W = Inches(13.333)
H = Inches(7.5)
OUT = "Agentic_AI_Brief.pptx"
TOTAL = 12
VERSION = "v1.7.0"


def set_run(run, size=18, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color=PAPER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def bar(slide, y=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, y, W, Inches(0.12))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ORANGE
    shape.line.fill.background()


def star(slide, left, top, size):
    s = slide.shapes.add_shape(MSO_SHAPE.STAR_4_POINT, left, top, size, size)
    s.fill.solid()
    s.fill.fore_color.rgb = ORANGE
    s.line.fill.background()


def textbox(slide, l, t, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)


def bullets(slide, l, t, w, h, items, size=18):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = "•  " + item
        set_run(run, size=size, color=INK)


def card(slide, l, t, w, h):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = WHITE
    s.line.color.rgb = LINE
    s.adjustments[0] = 0.08


def footer(slide, page):
    textbox(slide, Inches(0.6), Inches(7.12), Inches(9), Inches(0.3), f"Agentic AI  ·  {VERSION}  ·  Abhishek Mishra", 12, False, MUTED)
    textbox(slide, Inches(11.2), Inches(7.12), Inches(1.6), Inches(0.3), f"{page} / {TOTAL}", 12, False, MUTED, PP_ALIGN.RIGHT)


def content(prs, title, items, page, size=18):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    bar(s)
    textbox(s, Inches(0.7), Inches(0.4), Inches(12), Inches(0.7), title, 30, True, INK, font="Georgia")
    bullets(s, Inches(0.85), Inches(1.35), Inches(11.5), Inches(5.4), items, size)
    footer(s, page)


def cards3(prs, title, triples, page):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    bar(s)
    textbox(s, Inches(0.7), Inches(0.4), Inches(12), Inches(0.7), title, 30, True, INK, font="Georgia")
    gap = Inches(0.28)
    left = Inches(0.7)
    cw = Inches(3.85)
    for i, (h, body) in enumerate(triples):
        x = left + i * (cw + gap)
        card(s, x, Inches(1.4), cw, Inches(5.0))
        textbox(s, x + Inches(0.25), Inches(1.65), cw - Inches(0.45), Inches(0.7), h, 18, True, ORANGE_DARK)
        textbox(s, x + Inches(0.25), Inches(2.5), cw - Inches(0.45), Inches(3.6), body, 15, False, INK)
    footer(s, page)


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 Title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, CREAM)
    stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.28), H)
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = ORANGE
    stripe.line.fill.background()
    star(s, Inches(0.85), Inches(1.6), Inches(0.55))
    textbox(s, Inches(0.85), Inches(2.35), Inches(11), Inches(0.35), "BRIEF PROJECT OVERVIEW", 14, True, ORANGE_DARK)
    textbox(s, Inches(0.85), Inches(2.8), Inches(11), Inches(0.9), "Agentic AI", 48, True, INK, font="Georgia")
    textbox(
        s,
        Inches(0.85),
        Inches(3.75),
        Inches(11),
        Inches(0.8),
        "A ReAct assistant website — Next.js UI, FastAPI backend, tools, files, and credit top-ups.",
        20,
        False,
        MUTED,
    )
    textbox(s, Inches(0.85), Inches(5.1), Inches(11), Inches(0.35), f"{VERSION}  ·  Built by Abhishek Mishra", 16, True, INK)
    textbox(s, Inches(0.85), Inches(5.55), Inches(11), Inches(0.35), "Live: agentic-ai-d0cx.onrender.com", 15, False, MUTED)

    # 2 What it is
    content(
        prs,
        "What it is",
        [
            "Not a plain chatbot — an agent that thinks, calls tools, reads results, then answers.",
            "Web app: login, chat, voice, files/photos, Teach memory, Projects, Notes, Artifacts.",
            "Single agent for quick asks, or Researcher + Writer for longer research.",
            "Built to learn real agent design — the ReAct loop is visible Python, not a black box.",
        ],
        2,
    )

    # 3 Architecture
    cards3(
        prs,
        "Architecture (3 layers)",
        [
            ("Frontend", "Next.js 15 · React 19\n\nChat UI, Artifacts, Settings, SEO.\nPort 3000\nProxies /api → FastAPI"),
            ("Backend", "FastAPI · Uvicorn\n\nAuth, SSE chat stream, files, Teach, notes, billing.\nPort 7860"),
            ("Agent core", "Custom ReAct loop\n\nGroq LLM + tools + memory\nOptional crew\nCode sandbox"),
        ],
        3,
    )

    # 4 Stack
    content(
        prs,
        "Tech stack",
        [
            "LLM: Groq (openai/gpt-oss-120b). Vision: Qwen for photos.",
            "Python · FastAPI · Next.js · TypeScript · DuckDuckGo · wttr.in",
            "Deploy: one Docker image on Render (Vercel alone cannot run the Python API).",
            "UI: cream / paper / gold accent — Fraunces + Figtree.",
        ],
        4,
    )

    # 5 Features
    content(
        prs,
        "Key features",
        [
            "Chat + Code mode, Hindi / Bhojpuri, voice in/out, edit & retry.",
            "Attach PDF/Word/txt or up to 4 photos (gallery / camera).",
            "Teach + Projects = lasting instructions and files (memory, not GPU training).",
            "Landing-page builder → Artifacts preview + ZIP; AI images for pages.",
            "Tools: search, weather, calculator, notes, sandboxed Python, and more.",
        ],
        5,
        size=17,
    )

    # 6 How agent works
    content(
        prs,
        "How the agent works (ReAct)",
        [
            "Think → decide answer or tool.",
            "Act → call tool with JSON (search, weather, code_run…).",
            "Observe → read the tool result.",
            "Repeat (default max 8 steps) → clear human answer with sources.",
        ],
        6,
    )

    # 7 Limits
    content(
        prs,
        "Limits (at a glance)",
        [
            "Photos per message: max 4. Chat history sent to the model: last 16 turns.",
            "Free daily quota: 20 chats + 5 photo messages (then top-up).",
            "Recents chats: stored in browser localStorage (no fixed chat-count cap).",
            "Project files: up to 8. Teach training files: up to 5.",
        ],
        7,
    )

    # 8 Billing
    cards3(
        prs,
        "Credits & Razorpay top-ups",
        [
            ("Free", "₹0 / day\n\n20 chats\n5 photo messages\nResets every day"),
            ("Student", "₹49 one-time\n\n+50 chats\n+20 photo msgs\nFor learning / personal use"),
            ("Business", "₹199 one-time\n\n+200 chats\n+80 photo msgs\nFor shops & heavy use"),
        ],
        8,
    )

    # 9 Security
    content(
        prs,
        "Login & security",
        [
            "Name + email login, or Google / GitHub OAuth (httpOnly cookie session).",
            "API keys stay on the server (.env / Render env) — never shown in the UI.",
            "Code sandbox: no network, short timeout. Teach is prompt memory only.",
        ],
        9,
    )

    # 10 Run / demo
    content(
        prs,
        "Run locally & quick demo",
        [
            "API: python app.py → :7860   |   UI: cd frontend && npm run dev → :3000",
            "Demo: hello → Delhi weather → gold price → attach PDF → Teach facts → landing page → notes.",
            "Live Docker: set GROQ_API_KEY + RAZORPAY_KEY_ID / SECRET on Render.",
        ],
        10,
        size=17,
    )

    # 11 Scope
    content(
        prs,
        "What it is not",
        [
            "Not Claude Code / full computer use / OS control.",
            "Does not fine-tune a neural model on GPU.",
            "No Gmail/Drive/Calendar OAuth in v1.",
            "Credit top-ups are one-time packs — not a monthly subscription yet.",
        ],
        11,
    )

    # 12 Thanks
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, CREAM)
    stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.28), H)
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = ORANGE
    stripe.line.fill.background()
    star(s, Inches(0.85), Inches(2.0), Inches(0.55))
    textbox(s, Inches(0.85), Inches(2.8), Inches(11), Inches(0.8), "Thank you", 44, True, INK, font="Georgia")
    textbox(s, Inches(0.85), Inches(3.7), Inches(11), Inches(0.5), "Built by Abhishek Mishra", 22, True, ORANGE_DARK)
    textbox(s, Inches(0.85), Inches(4.4), Inches(11), Inches(0.5), "Questions? Open the app and try the demo flow.", 18, False, MUTED)
    textbox(s, Inches(0.85), Inches(5.3), Inches(11), Inches(0.35), "http://127.0.0.1:3000", 18, True, INK)
    textbox(s, Inches(0.85), Inches(5.75), Inches(11), Inches(0.35), "https://agentic-ai-d0cx.onrender.com", 16, False, MUTED)

    try:
        prs.save(OUT)
        print("saved:", OUT, "slides:", TOTAL)
    except PermissionError:
        alt = "Agentic_AI_Brief_v1.7.pptx"
        prs.save(alt)
        print("locked, saved:", alt)


if __name__ == "__main__":
    main()
